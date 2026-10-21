"""Gera e formata arquivos Excel com o historico de saidas e estornos."""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.relatorio_financeiro import valores_historicos


def gerar_excel(material_outputs):
    """Cria um relatorio Excel com quantidades e valores historicos."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Saídas"
    headers = [
        "Data/Hora",
        "Funcionário",
        "Projeto",
        "Material",
        "Qtd. original",
        "Unidade",
        "Qtd. estornada",
        "Qtd. líquida",
        "Valor unitário",
        "Valor original",
        "Valor estornado",
        "Valor líquido",
        "Situação",
    ]
    sheet.append(headers)
    historico = valores_historicos(material_outputs)

    for output in material_outputs:
        valores = historico.get(output.id)
        row = [
            output.created_at.strftime("%d/%m/%Y %H:%M:%S") if output.created_at else "",
            output.employee_name,
            output.project_name or "",
            output.material,
            float(output.quantity),
            valores["unit"] if valores else "",
            float(valores["returned_qty"]) if valores else None,
            float(valores["net_qty"]) if valores else None,
            float(valores["unit_price"]) if valores else None,
            float(valores["total_value"]) if valores else None,
            float(valores["returned_value"]) if valores else None,
            float(valores["net_value"]) if valores else None,
            valores["status"] if valores else "Legado sem preço",
        ]
        sheet.append(row)

        # Aplica formatos brasileiros sem alterar os valores exportados.
        sheet.cell(sheet.max_row, 9).number_format = '"R$" #,##0.000000'
        for column in (10, 11, 12):
            sheet.cell(sheet.max_row, column).number_format = '"R$" #,##0.00'
        for column in (5, 7, 8):
            sheet.cell(sheet.max_row, column).number_format = "#,##0.###"

    # Destaca o cabecalho e mantem as colunas legiveis ao abrir a planilha.
    header_fill = PatternFill(start_color="3C1276", end_color="3C1276", fill_type="solid")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [21, 24, 25, 24, 17, 12, 19, 17, 19, 19, 19, 19, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    file_buffer = BytesIO()
    workbook.save(file_buffer)
    file_buffer.seek(0)
    return file_buffer


def obter_nome_arquivo_excel():
    """Monta o nome do arquivo mensal de relatorio."""
    now = datetime.now()
    return f"saida_materiais_{now.year}_{now.month:02}.xlsx"
