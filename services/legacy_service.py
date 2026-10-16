from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from flask_login import current_user
from extensions import db
from models import Employee, Project, MaterialOutput
from services.relatorio_financeiro import valores_historicos

def get_nomes():
    return [item.name for item in Employee.query.order_by(Employee.name).all()]


def get_projetos():
    return [item.name for item in Project.query.order_by(Project.name).all()]


def get_report_suggestions():
    suggestion_map = {'projects': [], 'materials': [], 'employees': []}
    try:
        projects = db.session.query(MaterialOutput.project_name).filter(MaterialOutput.project_name.isnot(None)).distinct().order_by(MaterialOutput.project_name).all()
        suggestion_map['projects'] = [item[0] for item in projects if item[0]]

        materials = db.session.query(MaterialOutput.material).filter(MaterialOutput.material.isnot(None)).distinct().order_by(MaterialOutput.material).all()
        suggestion_map['materials'] = [item[0] for item in materials if item[0]]

        employees = db.session.query(MaterialOutput.employee_name).filter(MaterialOutput.employee_name.isnot(None)).distinct().order_by(MaterialOutput.employee_name).all()
        suggestion_map['employees'] = [item[0] for item in employees if item[0]]
    except Exception:
        db.session.rollback()
        raise
    return suggestion_map


def registrar_saida(dados):
    recorded_by = current_user.id if current_user.is_authenticated else None
    material_output = MaterialOutput(
        employee_name=dados[1],
        project_name=dados[2],
        material=dados[3],
        quantity=int(dados[4]),
        recorded_by=recorded_by,
    )
    db.session.add(material_output)
    db.session.commit()


def gerar_excel(material_outputs):
    wb=Workbook()
    ws=wb.active
    ws.title='Saídas'
    headers=['Data/Hora','Funcionário','Projeto','Material','Qtd. original','Unidade',
        'Qtd. estornada','Qtd. líquida','Valor unitário','Valor original','Valor estornado','Valor líquido','Situação']
    ws.append(headers)
    valores=valores_historicos(material_outputs)
    for o in material_outputs:
        v=valores.get(o.id)
        row=[o.created_at.strftime('%d/%m/%Y %H:%M:%S') if o.created_at else '',
            o.employee_name,o.project_name or '',o.material,float(o.quantity),
            v['unit'] if v else '',float(v['returned_qty']) if v else None,
            float(v['net_qty']) if v else None,float(v['unit_price']) if v else None,
            float(v['total_value']) if v else None,float(v['returned_value']) if v else None,
            float(v['net_value']) if v else None,v['status'] if v else 'Legado sem preço']
        ws.append(row)
        ws.cell(ws.max_row, 9).number_format='"R$" #,##0.000000'
        for col in (10,11,12):
            ws.cell(ws.max_row,col).number_format='"R$" #,##0.00'
        for col in (5,7,8):
            ws.cell(ws.max_row,col).number_format='#,##0.###'
    headfill=PatternFill(start_color='3C1276',end_color='3C1276',fill_type='solid')
    for cell in ws[1]:
        cell.font=Font(bold=True,color='FFFFFF')
        cell.fill=headfill
        cell.alignment=Alignment(horizontal='center',vertical='center')
    ws.freeze_panes='A2'
    ws.auto_filter.ref=ws.dimensions
    from openpyxl.utils import get_column_letter
    for i,width in enumerate([21,24,25,24,17,12,19,17,19,19,19,19,20],1):
        ws.column_dimensions[get_column_letter(i)].width=width
    output=BytesIO();wb.save(output);output.seek(0)
    return output


def obter_nome_arquivo_excel():
    agora = datetime.now()
    return f"saida_materiais_{agora.year}_{agora.month:02}.xlsx"


